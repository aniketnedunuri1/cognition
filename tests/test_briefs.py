import json
from datetime import date

import pytest

from conftest import load
from server import briefs, cache, exa, feedback

POST = ("Brex is migrating our monolith from Java 8 to Java 17 this quarter. Phil Burrows leads the platform "
        "team. Teams at Nubank saw an 8x engineering time efficiency gain.")


@pytest.fixture
def env(tmp_path, monkeypatch, stub):
    monkeypatch.setenv("BRIEFS_DIR", str(tmp_path / "briefs"))
    kd = tmp_path / "k"; kd.mkdir()
    (kd / "use_cases.json").write_text(json.dumps({"use_cases": [{"name": "Code migration & modernization"}]}))
    (kd / "customers.json").write_text(json.dumps({"customers": []}))
    monkeypatch.setenv("KNOWLEDGE_DIR", str(kd))
    blog = cache.upsert_doc(cache.make_doc(identity="https://brex.com/blog/java", kind="web",
                                           url="https://brex.com/blog/java", title="Java 17", text=POST,
                                           published_date=f"{date.today()}T00:00:00Z",
                                           retrieved_at=cache.now_iso()))["source_id"]
    stub("server.exa._post", lambda path, body: load("exa_people_brex.json"))
    phil = next(p["source_id"] for p in exa.search_people("Brex", ["platform"])["people"] if p["name"] == "Phil Burrows")
    return blog, phil


EMAIL = ("Saw the Brex platform team is migrating the monolith from Java 8 to Java 17 this quarter. Teams at "
         "Nubank used Devin on a migration like that and saw an 8x engineering time efficiency gain. Devin takes "
         "the repetitive upgrade work in parallel and opens PRs your engineers review, so the platform team keeps "
         "shipping while the migration moves. Would twenty minutes next week be useful to walk through how that "
         "would look for your services? Happy to send a short write-up first if that is easier for you.")
NOTE = "Phil, saw the platform team's Java 17 migration. Teams at Nubank used Devin on similar work. Would love to connect."
FOLLOW = ("Thanks for connecting, Phil. One thing I did not put in my note: on Java upgrades, Devin works through the "
          "repetitive changes across services in parallel and opens a PR per service with tests run, so reviewers "
          "see small diffs instead of one giant change. If it helps, I can share how teams structure the rollout so "
          "the platform team stays in control of what merges and when. Open to a short call?")


def brief(blog, phil, *, with_hook=True, signal=True):
    ev = [{"source_id": blog, "quote": "migrating our monolith from Java 8 to Java 17"}]
    person = {"name": "Phil Burrows", "title": "Head of Platform & Infra Engineering", "deal_role": "champion",
              "why": "Leads platform", "evidence": [{"source_id": phil}],
              "hooks": [{"value": "Leads the platform team", "evidence": [{"source_id": blog, "quote": "Phil Burrows leads the platform team"}]}] if with_hook else []}
    return {"company_name": "Brex",
            "verdict": {"verdict": "act_now", "why_now": "Java 17 migration underway", "based_on": ev},
            "signals": [{"headline": "Java 8 → 17 migration this quarter", "type": "migration", "horizon": "tactical",
                         "evidence": ev, "use_case": "Code migration & modernization", "strength": "strong",
                         "who_cares": []}] if signal else [],
            "people": [person],
            "top_contacts": [{"name": "Phil Burrows", "why": {"text": "Owns the migration", "based_on": [{"source_id": phil}]}}],
            "drafts": [{"name": "Phil Burrows", "email": {"subject": "Your Java 17 migration", "body": EMAIL},
                        "linkedin_note": NOTE, "linkedin_followup": FOLLOW, "built_from": ev + [{"source_id": blog}],
                        "claims": [{"text": "Brex is migrating its monolith to Java 17", "source_id": blog,
                                    "quote": "migrating our monolith from Java 8 to Java 17"}]}]}


def test_save_render_and_slack(env):
    blog, phil = env
    out = briefs.save_brief("https://www.brex.com", brief(blog, phil))
    assert out["report"]["dropped_count"] == 0, out["report"]
    rec = json.loads(next((briefs.bdir() / "brex.com").glob("*.json")).read_text())
    assert rec["brief"]["drafts"][0]["low_confidence"] is False
    md = (briefs.bdir() / "brex.com" / f"{rec['date']}.md").read_text()
    assert "https://brex.com/blog/java" in md and "[LinkedIn](https://linkedin.com/in/phil-burrows)" in md
    assert "Act now" in out["slack_summary"] and "https://brex.com/blog/java" in out["slack_summary"]
    assert briefs.get_last_brief("brex.com")["found"]


def test_low_confidence_label_set_by_validator(env):
    blog, phil = env
    briefs.save_brief("brex.com", brief(blog, phil, with_hook=False, signal=False))
    rec = briefs.load_record("brex.com")[1]
    assert rec["brief"]["drafts"][0]["low_confidence"] is True
    assert briefs.LOW_CONFIDENCE in briefs.render_markdown(rec)
    assert any("act_now needs" in f["flag"] for f in rec["report"]["flags"])


def test_draft_with_blocking_lint_dropped(env):
    blog, phil = env
    b = brief(blog, phil)
    b["drafts"][0]["linkedin_note"] = "Phil, Devin and Windsurf cut migrations 50x. Connect?"
    out = briefs.save_brief("brex.com", b)
    assert out["report"]["dropped"][0]["path"] == "drafts[0]"
    reason = out["report"]["dropped"][0]["reason"]
    assert "Windsurf" in reason and "50x" in reason


def test_since_last_run_diff(env):
    blog, phil = env
    old = briefs.validate_brief({"domain": "brex.com", **brief(blog, phil, signal=False)}, use_cases=None)[0]
    new = briefs.validate_brief({"domain": "brex.com", **brief(blog, phil)}, use_cases=None)[0]
    d = briefs.diff_briefs(old, new)
    assert d["new_signals"] == ["Java 8 → 17 migration this quarter"] and not d["nothing_new"]
    assert briefs.diff_briefs(new, new)["nothing_new"]


def test_feedback_rules_and_examples(tmp_path, monkeypatch):
    monkeypatch.setenv("PLAYBOOK_DIR", str(tmp_path))
    assert not feedback.add_playbook_rule("x", "persona:ceo")["ok"]
    r = feedback.add_playbook_rule("Lead with the peer metric", "persona:economic_buyer")
    assert r["confirmation"] == "Learned: Lead with the peer metric. Applies to: persona:economic_buyer."
    assert feedback.add_playbook_rule("lead with the peer metric!", "persona:economic_buyer")["evidence"] == 2
    r = feedback.add_playbook_rule("Lead with their signal", "persona:economic_buyer", replaces="Lead with the peer metric")
    assert r["replaced"] == "Lead with the peer metric"
    assert [x["rule"] for x in feedback.get_playbook()["learned_rules"]] == ["Lead with their signal"]
    assert "Lead with their signal" in (tmp_path / "learned_rules.md").read_text()
    for i in range(5):
        feedback.add_golden_example("champion", "email", f"m{i}", subject="s")
    assert [e["message"] for e in feedback.get_playbook()["golden_examples"]["champion/email"]] == ["m4", "m3", "m2"]



def _docx_text(path):
    from docx import Document
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    def cells(table):
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(p.text for p in cell.paragraphs)
                for nested in cell.tables:
                    cells(nested)
    for table in doc.tables:
        cells(table)
    return "\n".join(chunks)


def test_docx_and_download_link(env, monkeypatch):
    from starlette.testclient import TestClient
    from server.app import build_server
    blog, phil = env
    out = briefs.save_brief("brex.com", brief(blog, phil))
    stem = out["saved"].removesuffix(".json")
    assert out["slack_summary"].endswith(f"Full brief: {stem}.docx")  # CLI: local path
    text = _docx_text(f"{stem}.docx")
    assert "Brex — account brief" in text and "Claims to check" in text and "Your Java 17 migration" in text
    assert "Who to work" in text

    monkeypatch.setattr(briefs, "BASE_URL", "https://tunnel.example")
    link = briefs.brief_link(briefs.load_record("brex.com")[1])
    assert link.startswith("https://tunnel.example/briefs/") and link.endswith("/brex.com-" + stem[-10:] + ".docx")
    briefs.save_brief("brex.com", brief(blog, phil))  # a re-save keeps the link
    assert briefs.brief_link(briefs.load_record("brex.com")[1]) == link

    with TestClient(build_server("secret").http_app()) as client:
        r = client.get(link.removeprefix("https://tunnel.example"))
        assert r.status_code == 200 and r.content[:2] == b"PK"  # a .docx, no bearer token needed
        assert client.get("/briefs/not-a-token/brex.com.docx").status_code == 404
        assert client.post("/mcp", json={}).status_code == 401  # the MCP endpoint still needs the token


def test_render_fixes():
    from server.briefs import _local_day, _named_amounts, _one_line, _plain, conflicting
    assert conflicting([{"value": "Apollo reports ~1,800 employees"}, {"value": "Tracxn reports 5,158 employees"}])
    assert not conflicting([{"value": "Apollo reports ~1,800 employees"}, {"value": "Revelio reports ~1,757 employees"}])
    assert conflicting([{"value": "Apollo reports $700M"}, {"value": "RocketReach reports $279M"}])
    assert not conflicting([{"value": "Kotlin on Micronaut"}, {"value": "Buildkite and AWS"}])
    assert _one_line("Building agents\n\nBy Graham Fuller") == "Building agents By Graham Fuller"
    assert _plain("Assessment: CTO first") == "CTO first" and _plain("CTO first") == "CTO first"
    assert _local_day("2026-10-08T00:30:00+00:00") in ("2026-10-07", "2026-10-08")  # depends on the machine's zone
    assert _named_amounts("Litmus reports the Capital One deal closed April 7 for $5.15B; PitchBook lists $4.52B") == [
        "Litmus $5.15B", "PitchBook $4.52B"]
    assert _named_amounts("Tracxn reports 5,158 employees as of Aug 31, 2026, and TipRanks reports 5,389.") == [
        "Tracxn 5,158", "TipRanks 5,389"]


def test_rendered_draft_shows_claims_and_no_false_conflicts(env):
    blog, phil = env
    b = brief(blog, phil)
    b["engineering"] = {"tech_stack": [{"value": "Kotlin", "evidence": [{"source_id": blog}]},
                                       {"value": "Java", "evidence": [{"source_id": blog}]}]}
    md = briefs.save_brief("brex.com", b)["brief_markdown"]
    assert "**Claims to check**" in md and '> "migrating our monolith from Java 8 to Java 17"' in md
    assert "sources differ" not in md
    assert "## Who to work" in md and "## Buying centers" not in md


def test_working_number_render(env):
    blog, phil = env
    b = brief(blog, phil)
    b["snapshot"] = {"headcount": [
        {"value": "Apollo reports ~1,800 employees", "evidence": [{"source_id": blog}]},
        {"value": "Revelio Labs reports ~1,757 employees", "evidence": [{"source_id": blog}]},
        {"value": "Tracxn reports 5,158 employees", "evidence": [{"source_id": blog}]},
    ], "revenue": [
        {"value": "Apollo reports $700M", "evidence": [{"source_id": blog}]},
        {"value": "RocketReach reports $279M", "evidence": [{"source_id": blog}]},
    ]}
    md = briefs.save_brief("brex.com", b)["brief_markdown"]
    head = next(l for l in md.splitlines() if l.startswith("| Headcount |"))
    assert "Apollo reports ~1,800 employees" in head
    assert "5,158" not in head and "1,757" not in head
    assert "sources differ" not in head
    assert "Other sources: Tracxn 5,158" in md
    assert "1,757" not in md
    rev = next(l for l in md.splitlines() if l.startswith("| Revenue |"))
    assert "$700M" in rev and "$279M" not in rev
    assert "RocketReach $279M" in md


def test_draft_notes_and_skeleton_flag(env):
    blog, phil = env
    b = brief(blog, phil)
    b["drafts"][0]["notes"] = "- Email only.\n- Do not pitch a product tour."
    md = briefs.save_brief("brex.com", b)["brief_markdown"]
    assert "**Send notes**" in md and "Email only." in md
    b["drafts"].append(dict(b["drafts"][0]))
    out = briefs.save_brief("brex.com", b)
    assert out["report"]["dropped_count"] == 0, out["report"]
    assert any("same opening" in f["flag"] for f in out["report"]["flags"])
