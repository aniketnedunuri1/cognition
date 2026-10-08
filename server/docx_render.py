"""The validated brief as a designed .docx (from the JSON record, not from Markdown)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.document import Document as DocType
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .briefs import (
    LOW_CONFIDENCE, ROLE, VERDICT, _Sources, _dated, _facts,
    _one_line, _plain, _url, working_cell, working_fact,
)
from .cache import get_doc

# US Letter content width at 0.75" margins: 8.5" - 1.5" = 7" = 10080 DXA.
CONTENT = 10080
INK = "1F2937"
MUTED = "6B7280"
LINE = "E5E7EB"
PAPER = "FFFFFF"
# Verdict (similarity: one color each).
VERDICT_FILL = {"act_now": "DCFCE7", "warm": "FEF3C7", "not_now": "F3F4F6"}
VERDICT_INK = {"act_now": "166534", "warm": "B45309", "not_now": "4B5563"}
# Deal role (reuse the same hue everywhere that role appears).
ROLE_INK = {
    "economic_buyer": "1E40AF", "champion": "0F766E", "technical_evaluator": "6D28D9",
    "gatekeeper": "4B5563", "parent_org_approver": "9A3412",
}
CHAMPION_INK = "047857"
INCUMBENT_INK = "B91C1C"
SIGNAL_INK = {"strong": "1D4ED8", "medium": "64748B"}
CARD = "F8FAFC"
DRAFT = "F8FAFC"
NOTES = "F3F4F6"
SIGNAL_FILL = "F8FAFC"

_XML_UNSAFE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _rgb(h: str) -> RGBColor:
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _font(run, *, size=11, bold=False, italic=False, color=INK, name="Arial"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = _rgb(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def _run(p: Paragraph, text: str, **kw) -> None:
    if not text:
        return
    _font(p.add_run(_XML_UNSAFE.sub("", text)), **kw)


def _link(p: Paragraph, text: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "Arial")
    props.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")  # 9pt
    props.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    props.append(u)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(props)
    run.append(t)
    link.append(run)
    p._p.append(link)


def _shade(cell: _Cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _borders(cell: _Cell, color: str | None = None, sz: str = "4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if color:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), sz)
            el.set(qn("w:color"), color)
            el.set(qn("w:space"), "0")
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    tcPr.append(tb)


def _cell_width(cell: _Cell, dxa: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(dxa))
    w.set(qn("w:type"), "dxa")
    tcPr.append(w)
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "80"), ("left", "120"), ("bottom", "80"), ("right", "120")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _tbl_width(table: Table, dxa: int, cols: list[int]) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.get_or_add_tblPr()
    w = tblPr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tblPr.append(w)
    w.set(qn("w:w"), str(dxa))
    w.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for el in list(grid):
            grid.remove(el)
        for c in cols:
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(c))
            grid.append(gc)
    look = tblPr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tblPr.append(look)
    look.set(qn("w:firstRow"), "0")
    # Keep the table from inheriting Table Grid.
    table.style = "Normal Table"


def _container(doc: Document, fill: str = CARD, width: int = CONTENT) -> _Cell:
    table = doc.add_table(rows=1, cols=1)
    _tbl_width(table, width, [width])
    cell = table.cell(0, 0)
    _cell_width(cell, width)
    _shade(cell, fill)
    _borders(cell, LINE, "4")
    cell.paragraphs[0].text = ""
    return cell


def _p(cell_or_doc, *, before=0, after=60) -> Paragraph:
    if isinstance(cell_or_doc, DocType):
        p = cell_or_doc.add_paragraph()
    else:
        first = cell_or_doc.paragraphs[0]
        if len(cell_or_doc.paragraphs) == 1 and not (first.text or "").strip():
            p = first
        else:
            p = cell_or_doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before / 20)
    pf.space_after = Pt(after / 20)
    pf.line_spacing = 1.12
    return p


def _cite(p: Paragraph, marker: str) -> None:
    if marker:
        _run(p, " " + marker, size=8, color=MUTED)


def brief_to_docx(rec: dict, path: Path) -> Path:
    b, src = rec["brief"], _Sources()
    doc = Document()
    doc.core_properties.title = f"{b['company_name']} account brief"
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(INK)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    _run(title, f"{b['company_name']} — account brief", size=20, bold=True)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(10)
    _run(meta, f"{rec['date']}  ·  {b['domain']}", size=9, color=MUTED)

    _verdict(doc, b, src)
    _snapshot(doc, b, src)
    _people(doc, b, src)
    _signals(doc, b, src)
    _drafts(doc, b, src)
    _abm(doc, b, src)
    _sources(doc, src, rec)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    _run(p, text, size=13, bold=True)


def _verdict(doc: Document, b: dict, src: _Sources) -> None:
    v = b.get("verdict")
    if not v:
        p = doc.add_paragraph()
        _run(p, "No supported verdict.", italic=True, color=MUTED)
        return
    key = v["verdict"]
    cell = _container(doc, VERDICT_FILL[key])
    p = _p(cell, after=40)
    _run(p, VERDICT[key], size=16, bold=True, color=VERDICT_INK[key])
    _run(p, "   ", size=11)
    _run(p, _plain(v["why_now"]), size=11)
    _cite(p, src.cite(v["based_on"]))
    if v.get("watch_for"):
        w = _p(cell, after=40)
        _run(w, "Watch for: ", size=9, italic=True, color=MUTED)
        _run(w, v["watch_for"], size=9, color=MUTED)


def _snapshot(doc: Document, b: dict, src: _Sources) -> None:
    _heading(doc, "Snapshot")
    snap, eng = b["snapshot"], b["engineering"]
    rows = [
        ("HQ", _facts(src, snap["hq"])),
        ("Parent", _facts(src, snap["parent_company"])),
    ]
    notes = []
    for key, lab in (("headcount", "Headcount"), ("revenue", "Revenue"), ("funding", "Funding")):
        cell, note = working_cell(src, snap[key])
        rows.append((lab, cell))
        if note:
            notes.append(f"{lab} — {note}")
    eng_hc = working_fact(eng["eng_headcount"])
    rows += [
        ("Eng size", _facts(src, [eng_hc] if eng_hc else [])),
        ("Stack", _facts(src, eng["tech_stack"][:1])),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    label_w, val_w = 1600, CONTENT - 1600
    _tbl_width(table, CONTENT, [label_w, val_w])
    for i, (lab, val) in enumerate(rows):
        lc, vc = table.cell(i, 0), table.cell(i, 1)
        _cell_width(lc, label_w)
        _cell_width(vc, val_w)
        _shade(lc, PAPER)
        _shade(vc, PAPER)
        for cell, color in ((lc, LINE), (vc, LINE)):
            _bottom_only(cell, color)
        lp, vp = lc.paragraphs[0], vc.paragraphs[0]
        lp.text = ""; vp.text = ""
        _run(lp, lab, size=9, bold=True, color=MUTED)
        _run(vp, val, size=10)
    for note in notes:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        _run(p, note, size=8, color=MUTED)


def _bottom_only(cell: _Cell, color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "right", "bottom"):
        el = OxmlElement(f"w:{edge}")
        if edge == "bottom":
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), color)
            el.set(qn("w:space"), "0")
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    tcPr.append(tb)


def _people(doc: Document, b: dict, src: _Sources) -> None:
    _heading(doc, "Who to work")
    if not b["people"]:
        p = doc.add_paragraph()
        _run(p, "Not found", italic=True, color=MUTED)
        return
    for p in b["people"]:
        cell = _container(doc, CARD)
        head = _p(cell, after=20)
        _run(head, p["name"], size=12, bold=True)
        _run(head, f"  {p['title']}", size=10, color=MUTED)
        badges = _p(cell, after=40)
        role = p["deal_role"]
        _run(badges, ROLE.get(role, role), size=8, bold=True, color=ROLE_INK.get(role, INK))
        if p["label"] == "likely_champion":
            _run(badges, "   Likely champion", size=8, bold=True, color=CHAMPION_INK)
        if p.get("incumbent_owner"):
            _run(badges, "   Incumbent owner", size=8, bold=True, color=INCUMBENT_INK)
        why = _p(cell, after=20)
        _run(why, _plain(p["why"]), size=10)
        _cite(why, src.cite(p["evidence"]))
        if url := _url(p.get("profile_source_id") or ""):
            _run(why, "  ", size=9)
            _link(why, "LinkedIn", url)
        hook = (p["hooks"] or [{}])[0].get("value") if p["hooks"] else None
        if hook:
            h = _p(cell, after=20)
            _run(h, hook, size=9, color=MUTED)
        # spacer after card
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        sp.paragraph_format.space_before = Pt(0)


def _signals(doc: Document, b: dict, src: _Sources) -> None:
    _heading(doc, "Intent signals")
    if not b["signals"]:
        p = doc.add_paragraph()
        _run(p, "No signals met the bar.", italic=True, color=MUTED)
    for s in sorted(b["signals"], key=lambda s: s["strength"] != "strong"):
        cell = _container(doc, SIGNAL_FILL)
        h = _p(cell, after=20)
        _run(h, s["headline"], size=11, bold=True)
        meta = _p(cell, after=40)
        _run(meta, s["strength"], size=8, bold=True, color=SIGNAL_INK[s["strength"]])
        _run(meta, f"  ·  {s['use_case']}", size=9, color=MUTED)
        if s["who_cares"]:
            _run(meta, f"  ·  Who cares: {', '.join(s['who_cares'])}", size=9, color=MUTED)
        for e in s["evidence"]:
            if not e.get("quote"):
                continue
            q = _p(cell, after=20)
            _run(q, f'"{_one_line(e["quote"])}"', size=9, italic=True, color=MUTED)
            _cite(q, src.cite([e]) + _dated(e))
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
    if b["competitive_context"]:
        _heading(doc, "Competitive context")
        for c in b["competitive_context"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            _run(p, c["value"], size=10)
            _run(p, f"  ({c['situation'].replace('_', ' ')})", size=9, color=MUTED)
            _cite(p, src.cite(c["evidence"]))


def _drafts(doc: Document, b: dict, src: _Sources) -> None:
    _heading(doc, "Drafted messages")
    p = doc.add_paragraph()
    _run(p, "Drafts for rep review. Never auto-sent.", size=9, italic=True, color=MUTED)
    people = {x["name"]: x for x in b["people"]}
    for d in b["drafts"]:
        cell = _container(doc, DRAFT)
        head = _p(cell, after=40)
        _run(head, d["name"], size=13, bold=True)
        person = people.get(d["name"])
        if person:
            role = person["deal_role"]
            _run(head, f"   {ROLE.get(role, role)}", size=8, bold=True, color=ROLE_INK.get(role, INK))
        if d["low_confidence"]:
            w = _p(cell, after=40)
            _run(w, LOW_CONFIDENCE, size=9, italic=True, color="B45309")
        sub = _p(cell, after=20)
        _run(sub, "Email — ", size=9, bold=True, color=MUTED)
        _run(sub, d["email"]["subject"], size=11, bold=True)
        for line in d["email"]["body"].splitlines() or [""]:
            bp = _p(cell, after=40)
            _run(bp, line or " ", size=11)
        ln = _p(cell, after=20)
        _run(ln, "LinkedIn connection note", size=9, bold=True, color=MUTED)
        _run(_p(cell, after=40), d["linkedin_note"], size=10)
        fu = _p(cell, after=20)
        _run(fu, "LinkedIn follow-up", size=9, bold=True, color=MUTED)
        for line in d["linkedin_followup"].splitlines() or [""]:
            _run(_p(cell, after=40), line or " ", size=10)
        if d.get("notes"):
            notes_p = _p(cell, after=20)
            _run(notes_p, "Send notes  ·  rep only, never sent", size=8, bold=True, color=MUTED)
            for line in d["notes"].splitlines():
                np = _p(cell, after=10)
                _run(np, line, size=9, color=MUTED)
        claims = _p(cell, after=20)
        _run(claims, "Claims to check", size=8, bold=True, color=MUTED)
        for c in d["claims"]:
            cp = _p(cell, after=8)
            _run(cp, _one_line(c["text"]), size=8, color=MUTED)
            _cite(cp, src.cite([c]))
            qp = _p(cell, after=16)
            _run(qp, f'"{_one_line(c["quote"])}"', size=8, italic=True, color=MUTED)
        built = _p(cell, after=20)
        _run(built, "Built from: ", size=8, color=MUTED)
        _cite(built, src.cite(d["built_from"]))
        for x in d["lint"]:
            lp = _p(cell, after=8)
            _run(lp, f"Lint ({x['channel']}): {x['issue']}", size=8, color=MUTED)
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(10)


def _abm(doc: Document, b: dict, src: _Sources) -> None:
    _heading(doc, "ABM themes")
    if not b["abm_themes"]:
        p = doc.add_paragraph()
        _run(p, "None.", italic=True, color=MUTED)
        return
    for a in b["abm_themes"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _run(p, _plain(a["text"]), size=10)
        _cite(p, src.cite(a["based_on"]))


def _sources(doc: Document, src: _Sources, rec: dict) -> None:
    _heading(doc, "Sources")
    for n, sid in enumerate(src.order, 1):
        d = get_doc(sid) or {}
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _run(p, f"{n}. ", size=8, color=MUTED)
        title = d.get("title") or d.get("url") or d.get("label") or "Source"
        if d.get("url"):
            _link(p, title, d["url"])
        else:
            _run(p, title, size=8, color=MUTED)
        when = ", ".join(x for x in (
            f"published {d['published_date'][:10]}" if d.get("published_date") else "",
            f"retrieved {d['retrieved_at'][:10]}" if d.get("retrieved_at") else "") if x)
        if when:
            _run(p, f"  — {when}", size=8, color=MUTED)
    rep = rec["report"]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    _run(p, f"Validator: {rep['dropped_count']} unsupported item(s) dropped, {len(rep['flags'])} flag(s).",
         size=8, italic=True, color=MUTED)
