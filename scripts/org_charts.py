#!/usr/bin/env python3
"""Sumble org charts → .docx for SDR multithreading.

For each account domain:
  1. Resolve the organization (POST /v9/organizations, match mode).
  2. Find senior engineering leaders (POST /v9/people, filter mode,
     job_level_min 'Head', ordered by seniority).
  3. Pick 3-5 decision-makers, then fetch their inferred managers and
     direct reports (POST /v9/people, list mode, related_people).
  4. Render an indented org chart per account into one .docx.

Usage:  python scripts/org_charts.py brex.com [acme.com ...] [--out org_charts.docx]
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import time
from pathlib import Path

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401  (kept for parity with renderer)
from docx.shared import Inches, Pt, RGBColor
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / ".env")

API = "https://api.sumble.com/v9"
MAX_DMS = 5
MAX_PER_ACCOUNT = 25
MAX_REPORTS_PER_DM = 8
MIN_CONFIDENCE = 0.4
LEADERS_LIMIT = 40

INK = "1F2937"
MUTED = "6B7280"
CARD = "F8FAFC"
LINE = "E5E7EB"
DM_INK = "1E40AF"

IRRELEVANT = re.compile(
    r"sales|marketing|financ|accounting|legal|people|talent|recruit|hr\b|"
    r"revenue|customer|support|success|operations|design|product\b|growth|"
    r"communications|counsel|compliance|facilities|\boffice\b|credit|risk|"
    r"\bgtm\b|banking|commercial|deal desk",
    re.I,
)
# Corporate execs who aren't an eng buying-committee member at larger orgs.
# At small companies (<SMALL_ORG employees) the CEO/founder often *is* the buyer.
NON_BUYER_EXEC = re.compile(r"president|founder|chief executive|\bceo\b|\bcco\b|\bcoo\b", re.I)
SMALL_ORG = 200
# Titles that make someone a plausible Devin decision-maker.
BUYER = re.compile(
    r"chief technology|cto\b|chief ai|cio\b|chief information|"
    r"(svp|vp|vice president).*(eng|tech|software|platform|infra|ai)|"
    r"eng.*(svp|vp|vice president)",
    re.I,
)
PLATFORM = re.compile(
    r"(head|director|senior director|manager).*(platform|developer|devprod|"
    r"devex|productivity|infra|eng|ai|tools|enablement)|"
    r"(platform|developer|devprod|devex|productivity|infra|ai).*(head|director)",
    re.I,
)
GATEKEEPER = re.compile(r"ciso|security|procurement", re.I)

LEVEL_RANK = {
    "CXO": 15, "SVP": 14, "VP": 13, "Executive Director": 12,
    "Senior Director": 11, "Director": 10, "General Manager": 9, "Head": 8,
    "Associate Director": 7, "Senior Manager": 6, "Manager": 5,
    "Principal": 4, "Lead": 3, "Senior": 2, "Individual Contributor": 1,
}


def _post(client: httpx.Client, path: str, body: dict) -> dict:
    r = client.post(f"{API}{path}", json=body, timeout=60)
    if r.status_code >= 400:
        raise RuntimeError(f"{path} {r.status_code}: {r.text[:400]}")
    return r.json()


def _poll_people(client: httpx.Client, body: dict) -> dict:
    """POST /people is async: start, then poll request_id until terminal."""
    res = _post(client, "/people", body)
    while res.get("status") in ("pending", "running"):
        time.sleep(2)
        res = _post(client, "/people", {"request_id": res["request_id"]})
    if res.get("status") == "failed":
        raise RuntimeError(f"people request failed: {res}")
    return res


def resolve_org(client: httpx.Client, domain: str) -> dict | None:
    res = _post(client, "/organizations", {
        "organizations": [{"url": domain}],
        "select": {"attributes": ["id", "name", "employee_count", "industry"]},
    })
    orgs = res.get("organizations") or []
    return orgs[0] if orgs else None


def senior_leaders(client: httpx.Client, org_id: int) -> tuple[list[dict], int]:
    res = _poll_people(client, {
        "filter": {
            "organization_ids": [org_id],
            "query": {"query": "job_function EQ 'Engineering & R&D' "
                               "AND job_level_min EQ 'Head'"},
        },
        "select": {"attributes": [
            "name", "job_title", "job_level", "linkedin_url", "location",
        ]},
        "order_by_column": "job_level",
        "limit": LEADERS_LIMIT,
    })
    return res.get("people") or [], res.get("total") or 0


def is_eng_relevant(title: str, *, allow_execs: bool = False) -> bool:
    if not title or IRRELEVANT.search(title):
        return False
    if not allow_execs and NON_BUYER_EXEC.search(title):
        return False
    return True


def deal_role(title: str) -> str:
    if GATEKEEPER.search(title):
        return "gatekeeper"
    if BUYER.search(title) or NON_BUYER_EXEC.search(title):
        return "economic_buyer"
    if PLATFORM.search(title):
        return "technical_evaluator"
    return "influencer"


def why_line(title: str, role: str) -> str:
    if role == "economic_buyer":
        return "Likely economic buyer — owns engineering budget and tooling decisions."
    if role == "technical_evaluator":
        return "Technical evaluator — owns the platform/dev-productivity surface Devin lands on."
    if role == "gatekeeper":
        return "Gatekeeper — will review security/procurement. Named, not a first touch."
    return "Senior eng leader — likely in the room for a tooling decision."


def pick_decision_makers(people: list[dict], *, small_org: bool = False) -> list[dict]:
    """Dedupe by (name, title); rank buyers first, then evaluators, by seniority."""
    seen, cands = set(), []
    for p in people:
        a = p.get("attributes") or {}
        key = (a.get("name"), a.get("job_title"))
        if not a.get("name") or not a.get("job_title") or key in seen:
            continue
        seen.add(key)
        if not is_eng_relevant(a["job_title"], allow_execs=small_org):
            continue
        role = deal_role(a["job_title"])
        cands.append({
            "person_id": p["person_id"],
            "name": a["name"],
            "title": a["job_title"],
            "level": a.get("job_level"),
            "linkedin": a.get("linkedin_url"),
            "location": a.get("location"),
            "role": role,
            "rank": LEVEL_RANK.get(a.get("job_level") or "", 0),
        })
    order = {"economic_buyer": 0, "technical_evaluator": 1, "influencer": 2, "gatekeeper": 3}
    cands.sort(key=lambda c: (order[c["role"]], -c["rank"]))
    # Keep at least one economic buyer if present, then fill to MAX_DMS.
    picked = cands[:MAX_DMS]
    if picked and not any(c["role"] == "economic_buyer" for c in picked):
        for c in cands:
            if c["role"] == "economic_buyer":
                picked[-1] = c
                break
    return picked


def related(client: httpx.Client, person_ids: list[int]) -> list[dict]:
    res = _poll_people(client, {
        "people": [{"person_id": pid} for pid in person_ids],
        "select": {
            "related_people": {
                "direction": ["managers", "direct_reports"],
                "attributes": ["name", "job_title", "job_level",
                               "linkedin_url", "confidence"],
            },
        },
    })
    return res.get("people") or []


def person_name(row: dict) -> str | None:
    return (row.get("attributes") or {}).get("name")


def account_chart(client: httpx.Client, domain: str) -> dict:
    org = resolve_org(client, domain)
    if not org or not (org.get("attributes") or {}).get("id"):
        return {"domain": domain, "error": "organization not found in Sumble"}
    oa = org["attributes"]
    leaders, total = senior_leaders(client, oa["id"])
    small_org = bool(oa.get("employee_count")) and oa["employee_count"] < SMALL_ORG
    dms = pick_decision_makers(leaders, small_org=small_org)
    seen_reports: set[str] = set()
    for dm, row in zip(dms, related(client, [d["person_id"] for d in dms])):
        rel = row.get("related_people") or {}
        mgrs = [m for m in rel.get("managers") or [] if person_name(m)]
        dm["manager"] = mgrs[0] if mgrs else None
        seen_names = {dm["name"]} | seen_reports
        reports = []
        for r in rel.get("direct_reports") or []:
            ra = r.get("attributes") or {}
            nm = ra.get("name")
            if not nm or not ra.get("job_title") or nm in seen_names:
                continue
            if not is_eng_relevant(ra["job_title"]):
                continue
            score = ((r.get("confidence") or {}).get("score")) or 0
            if score < MIN_CONFIDENCE:
                continue
            seen_names.add(nm)
            seen_reports.add(nm)
            reports.append({
                "name": nm,
                "title": ra["job_title"],
                "level": ra.get("job_level"),
                "linkedin": ra.get("linkedin_url"),
                "confidence": score,
                "rank": LEVEL_RANK.get(ra.get("job_level") or "", 0),
            })
        reports.sort(key=lambda r: (-r["confidence"], -r["rank"]))
        dm["reports"] = reports[:MAX_REPORTS_PER_DM]
    # Cap total rows (DMs + reports) at MAX_PER_ACCOUNT, trimming longest report lists.
    while sum(1 + len(d["reports"]) for d in dms) > MAX_PER_ACCOUNT:
        longest = max(dms, key=lambda d: len(d["reports"]))
        longest["reports"].pop()
        if not any(d["reports"] for d in dms):
            break
    return {
        "domain": domain,
        "name": oa.get("name") or domain,
        "employee_count": oa.get("employee_count"),
        "industry": oa.get("industry"),
        "leaders_found": total,
        "decision_makers": dms,
    }


# ---- docx ---------------------------------------------------------------


def _rgb(h: str) -> RGBColor:
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _run(p, text: str, *, size=10, bold=False, italic=False, color=INK):
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = _rgb(color)


def _link(p, text: str, url: str):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    props.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    props.append(u)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(props)
    run.append(t)
    link.append(run)
    p._p.append(link)


def render(charts: list[dict], out: Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Inches(0.75))
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(INK)

    title = doc.add_paragraph()
    _run(title, "Account org charts — Devin buying committee", size=16, bold=True)
    sub = doc.add_paragraph()
    _run(sub, "Decision-makers with inferred direct reports (Sumble). "
              "Reporting lines are inferred, not confirmed — the (0.x) on a "
              "report is Sumble's inference confidence.", size=9, color=MUTED)

    for i, acct in enumerate(charts):
        if i:
            doc.add_page_break()
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        _run(h, acct["name"], size=14, bold=True)
        meta = doc.add_paragraph()
        bits = [acct["domain"]]
        if acct.get("industry"):
            bits.append(acct["industry"])
        if acct.get("employee_count"):
            bits.append(f"~{acct['employee_count']:,} employees (Sumble)")
        _run(meta, "  ·  ".join(bits), size=9, color=MUTED)

        if acct.get("error"):
            p = doc.add_paragraph()
            _run(p, f"Not found — {acct['error']}", italic=True, color=MUTED)
            continue

        dms = acct["decision_makers"]
        if not dms:
            p = doc.add_paragraph()
            _run(p, "No senior engineering leaders found.", italic=True, color=MUTED)
            continue
        if acct["leaders_found"] > LEADERS_LIMIT:
            note = doc.add_paragraph()
            _run(note, f"Sumble lists {acct['leaders_found']} leaders at Head+; "
                       f"top {LEADERS_LIMIT} by seniority considered.",
               size=8, italic=True, color=MUTED)

        for dm in dms:
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(10)
            head.paragraph_format.space_after = Pt(1)
            head.paragraph_format.keep_with_next = True
            _run(head, dm["name"], size=11, bold=True)
            _run(head, f"  —  {dm['title']}", size=10)
            _run(head, f"   [{dm['role'].replace('_', ' ')}]", size=8, bold=True, color=DM_INK)
            line = doc.add_paragraph()
            line.paragraph_format.space_after = Pt(1)
            line.paragraph_format.keep_with_next = True
            mgr = (dm.get("manager") or {}).get("attributes") or {}
            if mgr.get("name"):
                _run(line, f"Reports to: {mgr['name']}"
                           + (f", {mgr.get('job_title')}" if mgr.get("job_title") else ""),
                     size=9, color=MUTED)
                _run(line, "   ", size=9)
            if dm.get("location"):
                _run(line, dm["location"], size=9, color=MUTED)
                _run(line, "   ", size=9)
            if dm.get("linkedin"):
                _link(line, "LinkedIn", dm["linkedin"])
            wp = doc.add_paragraph()
            wp.paragraph_format.space_after = Pt(2)
            wp.paragraph_format.keep_with_next = bool(dm["reports"])
            _run(wp, "Why: " + why_line(dm["title"], dm["role"]), size=9, italic=True)

            reps = dm["reports"]
            if reps:
                lbl = doc.add_paragraph()
                lbl.paragraph_format.space_after = Pt(0)
                lbl.paragraph_format.keep_with_next = True
                _run(lbl, f"    Direct reports ({len(reps)}):", size=9, bold=True, color=MUTED)
            for r in reps:
                rp = doc.add_paragraph()
                rp.paragraph_format.space_after = Pt(0)
                rp.paragraph_format.left_indent = Inches(0.35)
                _run(rp, f"•  {r['name']}", size=9)
                _run(rp, f" — {r['title']}", size=9, color=MUTED)
                if r.get("confidence"):
                    _run(rp, f"  ({r['confidence']:.1f})", size=8, color=MUTED)
                if r.get("linkedin"):
                    _run(rp, "  ", size=9)
                    _link(rp, "LinkedIn", r["linkedin"])

    doc.save(str(out))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("domains", nargs="+", help="Account domains, e.g. brex.com")
    ap.add_argument("--out", default="org_charts.docx", help="Output .docx path")
    args = ap.parse_args()

    key = os.environ.get("SUMBLE_API_KEY")
    if not key:
        sys.exit("SUMBLE_API_KEY not set (check .env)")

    charts = []
    with httpx.Client(headers={"Authorization": f"Bearer {key}"}) as client:
        for domain in args.domains:
            print(f"→ {domain} …", flush=True)
            try:
                charts.append(account_chart(client, domain.strip().lower()))
            except Exception as e:  # keep going for other accounts
                charts.append({"domain": domain, "name": domain, "error": str(e)})
                print(f"  ! {e}", file=sys.stderr)
            n = sum(1 + len(d["reports"]) for d in charts[-1].get("decision_makers", []))
            print(f"  {len(charts[-1].get('decision_makers', []))} decision-makers, {n} people total")

    out = render(charts, Path(args.out))
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
